"""
Converts Dedris-family (Sambhota) .doc files to Unicode Tibetan .docx, using the trained
lookup table from build_dedris_map.py (resources/dedris-map.json by default).

Since a single Dedris keystroke represents an entire pre-rendered consonant stack (see
DEDRIS_CONVERSION_PLAN.md / tibtexts/dedrismap.py), conversion is: resolve every (font, char)
keystroke independently via DedrisMap, concatenate the results in order, then run the same
punctuation cleanup insert_milestones.py's UniVol already uses on OCR/Sambhota-derived text.

Any keystroke the table has no confident answer for is left as an inline marker
(<<font:char>>) rather than guessed, and logged, so a human proofreader has an exact,
greppable list of spots to check. The output feeds unchanged into the existing
add_styles2docx.py step, same as the current Windows/udp.exe pipeline's .docx output does.

Usage:
    python convert_dedris.py KAMA-084-e.doc KAMA-084-f.doc
    python convert_dedris.py -w /other/workspace --map resources/dedris-map.json *.doc
"""
import argparse
import datetime
from glob import glob
from os import makedirs
from os.path import join, dirname, abspath, basename, splitext, exists

import docx

from tibtexts.fodtdoc import FodtDoc
from tibtexts.dedrismap import DedrisMap
from tibtexts.univol import UniVol

HERE = dirname(abspath(__file__))
DEFAULT_WORKSPACE = join(HERE, 'workspace')
DEFAULT_MAP_PATH = join(HERE, 'resources', 'dedris-map.json')

MIN_CONFIDENCE = 0.5  # below this, treat as unresolved rather than trust the guess


def convert_paragraph(paragraph, dmap, unresolved_log, para_index):
    out = []
    for char_index, (font, ch) in enumerate(paragraph):
        unicode_str, confidence = dmap.resolve(font, ch, min_confidence=MIN_CONFIDENCE)
        if unicode_str is None:
            out.append('<<{}:{}>>'.format(font, ch))
            unresolved_log.append((para_index, char_index, font, ch))
        else:
            out.append(unicode_str)
    text = ''.join(out)
    for pair in UniVol.normalize_pairs:
        text = text.replace(*pair)
    return text


def convert_doc(doc_path, dmap, workdir):
    fdoc = FodtDoc.from_doc(doc_path, workdir)
    unresolved_log = []
    paragraphs = [
        convert_paragraph(para, dmap, unresolved_log, i)
        for i, para in enumerate(fdoc.paragraphs)
    ]
    return paragraphs, unresolved_log


def write_docx(paragraphs, out_path):
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    d.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('docs', nargs='+', help='Sambhota .doc file(s) to convert, or a glob')
    parser.add_argument('-w', '--workspace', default=DEFAULT_WORKSPACE,
                         help='Workspace root (default: ./workspace next to this script)')
    parser.add_argument('--map', default=DEFAULT_MAP_PATH, help='Trained map JSON to use')
    parser.add_argument('-o', '--outdir', default=None,
                         help='Output directory for converted .docx (default: <workspace>/out)')
    args = parser.parse_args()

    dmap = DedrisMap.load(args.map)
    outdir = args.outdir or join(args.workspace, 'out')
    makedirs(outdir, exist_ok=True)
    logdir = join(args.workspace, 'logs')
    makedirs(logdir, exist_ok=True)
    temp_dir = join(args.workspace, 'temp')
    makedirs(temp_dir, exist_ok=True)

    doc_paths = []
    for pattern in args.docs:
        matches = glob(pattern)
        doc_paths.extend(matches if matches else [pattern])
    if not doc_paths:
        raise SystemExit("No input files found for {}".format(args.docs))

    timestamp = datetime.datetime.now().strftime('%Y-%m-%d_%H-%M')
    for doc_path in doc_paths:
        if not exists(doc_path):
            print("Skipping missing file: {}".format(doc_path))
            continue
        stem = splitext(basename(doc_path))[0]
        print("Converting {}...".format(doc_path))
        paragraphs, unresolved_log = convert_doc(doc_path, dmap, temp_dir)

        out_path = join(outdir, stem + '.docx')
        write_docx(paragraphs, out_path)

        fonts = dmap.fonts_encountered()
        total_chars = sum(len(p) for p in paragraphs)
        print("  Wrote {} ({} paragraphs, {} chars)".format(out_path, len(paragraphs), total_chars))
        print("  Fonts encountered so far (cumulative across this run): {}".format(', '.join(fonts)))
        if unresolved_log:
            print("  {} unresolved keystroke(s) -- see log".format(len(unresolved_log)))

        log_path = join(logdir, '{}-dedris-convert-{}.log'.format(stem, timestamp))
        with open(log_path, 'w', encoding='utf-8') as fout:
            fout.write("Dedris conversion log for {}\n".format(doc_path))
            fout.write("Fonts encountered: {}\n".format(', '.join(fonts)))
            fout.write("{} unresolved keystroke(s):\n".format(len(unresolved_log)))
            for para_index, char_index, font, ch in unresolved_log:
                fout.write("  paragraph {}, char {}: ({!r}, {!r})\n".format(
                    para_index, char_index, font, ch))
        print("  Log written to {}".format(log_path))


if __name__ == '__main__':
    main()
